"""Convert the canonical ODI Tracker workbook into dashboard input JSON."""

from __future__ import annotations

import argparse
import json
import os
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REQUIRED_COLUMNS = (
    "Investor/Company Name",
    "Announcement Date",
    "Announcement date (year-month)",
    "Destination Country/Territory",
    "Region",
    "Subregion",
    "Project Name",
    "Primary Investment/Contract Value",
    "Primary Industry",
    "Secondary Industry",
    "Primary Category",
    "Secondary Category",
    "Note",
    "Source Webpage URL",
    "Status",
)

VALID_PRIMARY_CATEGORIES = {"Direct Investment", "Contracted Construction"}


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def iso_date(value: Any, row_number: int) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    raw = text(value)
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(raw, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    raise ValueError(f"Projects row {row_number}: invalid Announcement Date {value!r}")


def numeric_or_none(value: Any, row_number: int) -> float | int | None:
    if value in (None, ""):
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(
            f"Projects row {row_number}: Primary Investment/Contract Value must be numeric or blank"
        )
    if value < 0:
        raise ValueError(f"Projects row {row_number}: project value cannot be negative")
    return value


def build(workbook_path: Path, output_path: Path, report_path: Path) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "Projects" not in workbook.sheetnames:
        raise ValueError("Workbook does not contain a Projects worksheet")

    worksheet = workbook["Projects"]
    rows = worksheet.iter_rows(values_only=True)
    try:
        headers = [text(value) for value in next(rows)]
    except StopIteration as exc:
        raise ValueError("Projects worksheet is empty") from exc

    missing_columns = [name for name in REQUIRED_COLUMNS if name not in headers]
    if missing_columns:
        raise ValueError(f"Projects worksheet is missing columns: {missing_columns}")

    projects: list[dict[str, Any]] = []
    seen: dict[tuple[str, ...], int] = {}
    errors: list[str] = []

    for row_number, values in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in values):
            continue
        row = dict(zip(headers, values))
        try:
            announced = iso_date(row.get("Announcement Date"), row_number)
            month = announced[:7]
            stated_month = text(row.get("Announcement date (year-month)")).replace("/", "-")
            if stated_month != month:
                raise ValueError(
                    f"Projects row {row_number}: date month {month} does not match {stated_month!r}"
                )

            required_text = (
                "Investor/Company Name",
                "Destination Country/Territory",
                "Region",
                "Subregion",
                "Project Name",
                "Primary Industry",
                "Primary Category",
                "Secondary Category",
            )
            for column in required_text:
                if not text(row.get(column)):
                    raise ValueError(f"Projects row {row_number}: {column} is blank")

            primary_category = text(row.get("Primary Category"))
            if primary_category not in VALID_PRIMARY_CATEGORIES:
                raise ValueError(
                    f"Projects row {row_number}: unexpected Primary Category {primary_category!r}"
                )

            secondary_industry = text(row.get("Secondary Industry")) or None
            if primary_category == "Direct Investment" and secondary_industry is None:
                raise ValueError(
                    f"Projects row {row_number}: Direct Investment requires Secondary Industry"
                )

            duplicate_key = tuple(
                text(row.get(column)).casefold()
                for column in (
                    "Investor/Company Name",
                    "Announcement Date",
                    "Destination Country/Territory",
                    "Project Name",
                    "Primary Category",
                )
            )
            if duplicate_key in seen:
                raise ValueError(
                    f"Projects row {row_number}: duplicates Projects row {seen[duplicate_key]}"
                )
            seen[duplicate_key] = row_number

            projects.append(
                {
                    "id": len(projects) + 1,
                    "investor": text(row.get("Investor/Company Name")),
                    "homeProvince": text(row.get("Chinese Company Home Province")) or None,
                    "date": announced,
                    "month": month,
                    "country": text(row.get("Destination Country/Territory")),
                    "region": text(row.get("Region")),
                    "subregion": text(row.get("Subregion")),
                    "project": text(row.get("Project Name")),
                    "valueUsd": numeric_or_none(
                        row.get("Primary Investment/Contract Value"), row_number
                    ),
                    "secondaryValue": text(
                        row.get("Secondary Investment/Contract Value")
                    )
                    or None,
                    "secondaryValueNote": text(
                        row.get("Secondary Investment/Contract Value Note")
                    )
                    or None,
                    "primaryCategory": primary_category,
                    "secondaryCategory": text(row.get("Secondary Category")),
                    "primaryIndustry": text(row.get("Primary Industry")),
                    "secondaryIndustry": secondary_industry,
                    "note": text(row.get("Note")),
                    "status": text(row.get("Status")) or None,
                    "motivation": text(row.get("Motivation")) or None,
                    "sourceUrl": text(row.get("Source Webpage URL")) or None,
                }
            )
        except ValueError as exc:
            errors.append(str(exc))

    if errors:
        preview = "\n".join(f"- {item}" for item in errors[:25])
        suffix = "" if len(errors) <= 25 else f"\n- ...and {len(errors) - 25} more"
        raise ValueError(f"Workbook validation failed:\n{preview}{suffix}")
    if not projects:
        raise ValueError("Projects worksheet contains no project records")

    previous_count = None
    if output_path.exists():
        try:
            previous_count = len(json.loads(output_path.read_text(encoding="utf-8"))["projects"])
        except (KeyError, TypeError, ValueError, json.JSONDecodeError):
            previous_count = None
    if (
        previous_count
        and len(projects) < previous_count * 0.9
        and os.environ.get("ALLOW_LARGE_PROJECT_DROP") != "1"
    ):
        raise ValueError(
            f"Project count fell from {previous_count} to {len(projects)}. "
            "Set ALLOW_LARGE_PROJECT_DROP=1 only after reviewing the deletion."
        )

    months = sorted({project["month"] for project in projects})
    payload = {"sourceWorkbook": workbook_path.name, "projects": projects}
    report = {
        "sourceWorkbook": workbook_path.name,
        "projectCount": len(projects),
        "directInvestmentCount": sum(
            project["primaryCategory"] == "Direct Investment" for project in projects
        ),
        "contractedConstructionCount": sum(
            project["primaryCategory"] == "Contracted Construction" for project in projects
        ),
        "startMonth": months[0],
        "endMonth": months[-1],
        "disclosedPrimaryValueCount": sum(
            project["valueUsd"] is not None for project in projects
        ),
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8"
    )
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return report


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--workbook", type=Path, default=Path("data/ODI_Tracker_latest.xlsx")
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("dashboard/source/data/raw-projects.json"),
    )
    parser.add_argument(
        "--report",
        type=Path,
        default=Path("dashboard/source/data/build-report.json"),
    )
    args = parser.parse_args()
    print(json.dumps(build(args.workbook, args.output, args.report), indent=2))


if __name__ == "__main__":
    main()
